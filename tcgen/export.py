"""Module 9: JSON / CSV / Excel export. Works on the plain-dict result produced by the agent."""
from __future__ import annotations

import io
import json

import pandas as pd

COLUMNS = ["ID", "Scenario", "Preconditions", "Steps", "Test Data", "Expected Result",
           "Priority", "Test Type", "Positive/Negative", "Requirements"]


def cases_df(cases: list[dict], numbered_steps: bool = True) -> pd.DataFrame:
    rows = []
    for c in cases:
        steps = c.get("steps", [])
        steps_txt = "\n".join(f"{i}. {s}" for i, s in enumerate(steps, 1)) if numbered_steps else " → ".join(steps)
        rows.append([c.get("id", ""), c.get("scenario", ""), c.get("preconditions", ""), steps_txt,
                     c.get("test_data", ""), c.get("expected_result", ""), c.get("priority", ""),
                     c.get("test_type", ""), c.get("polarity", ""), ", ".join(c.get("requirement_ids", []))])
    return pd.DataFrame(rows, columns=COLUMNS)


def summary_df(result: dict) -> pd.DataFrame:
    m, c = result["metrics"], result["classification"]
    rows = [("Category", c["category"]), ("Story quality", f'{c["quality"]} ({c["quality_score"]:.0f}%)'),
            ("Classifier", c["source"]), ("Requirements", m["requirements"]), ("Total test cases", m["total"]),
            ("Positive", m["positive"]), ("Negative", m["negative"]),
            ("Requirement coverage %", m["coverage_percent"]), ("Relevance %", m["relevance_percent"]),
            ("Avg similarity %", m["avg_similarity_percent"]), ("Duplicate %", m["duplicate_percent"]),
            ("Refinement rounds", m["refinements"]), ("LLM", result.get("llm_status", ""))]
    return pd.DataFrame(rows, columns=["Metric", "Value"])


def to_json_str(result: dict) -> str:
    return json.dumps({k: result[k] for k in ("story", "classification", "analysis", "test_cases", "metrics")},
                      indent=2, ensure_ascii=False)


def to_csv_bytes(cases: list[dict]) -> bytes:
    return cases_df(cases).to_csv(index=False).encode("utf-8-sig")


def to_excel_bytes(result: dict) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as xw:
        cases_df(result["test_cases"]).to_excel(xw, sheet_name="Test Cases", index=False)
        summary_df(result).to_excel(xw, sheet_name="Summary", index=False)
        ws = xw.sheets["Test Cases"]
        widths = [8, 32, 32, 48, 34, 40, 10, 16, 16, 14]
        for col, w in zip("ABCDEFGHIJ", widths):
            ws.column_dimensions[col].width = w
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        xw.sheets["Summary"].column_dimensions["A"].width = 28
        xw.sheets["Summary"].column_dimensions["B"].width = 44
    return bio.getvalue()


def batch_excel_bytes(results: list[dict]) -> bytes:
    frames = []
    for i, r in enumerate(results, 1):
        df = cases_df(r["test_cases"])
        df.insert(0, "Story #", i)
        frames.append(df)
    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as xw:
        pd.concat(frames, ignore_index=True).to_excel(xw, sheet_name="All Test Cases", index=False)
        pd.DataFrame([{"Story #": i, "Story": r["story"], "Category": r["classification"]["category"],
                       "Test cases": r["metrics"]["total"], "Coverage %": r["metrics"]["coverage_percent"],
                       "Relevance %": r["metrics"]["relevance_percent"]} for i, r in enumerate(results, 1)]
                     ).to_excel(xw, sheet_name="Summary", index=False)
    return bio.getvalue()
